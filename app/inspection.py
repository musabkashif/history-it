from __future__ import annotations
import base64, json, time
from pathlib import Path
from app.credits import as_credit_error
from app.diagnostics import Diagnostics
from app.images import JPEG_QUALITY, ORIGINAL_MAX_BYTES, ORIGINAL_MAX_EDGE, RESIZE_MAX_EDGE, prepare_model_image
from app.models import InspectionOutput

OUTPUT_SCHEMA = """Structured output schema only. This text does not decide record boundaries, record types, templates, or metadata readiness. Those judgments come from the client runtime prompt, the governing documents, the authorized images, and any client-approved calibration example in the user payload.

Return Lean Split v2 Asset Staging and Record Staging as the structured object. Do not update a spreadsheet, choose a batch, or calculate cumulative totals. All values must be strings. Inspect every authorized image. Preserve every authorized filename exactly once, in manifest order. Each Assigned Record ID must equal a Proposed Record ID in this response. Record Kind is Normal Odyssey or Metadata Only. Record Operation is Create New or Update Existing Open Record. Copy context.prior_open_item from the payload when continuing that record; do not replace a None value and do not create a second record for it. Copy context.folder_code exactly.

When Boundary State After Record starts with Open, or Metadata Readiness is Review, write the visual evidence for that decision in Review Note. Leave Review Note blank only when the record is neither Review nor Open. A folder boundary does not by itself close an item."""

def _runtime_prompt(settings):
    configured = getattr(settings, 'runtime_prompt_path', None) or ''
    candidates = []
    if configured:
        candidates.append(Path(configured))
    app_dir = Path(__file__).resolve().parents[1]
    candidates.append(app_dir.parent / 'Runtime_prompt_old_chatgpt.md')
    candidates.append(app_dir / 'prompts' / 'Runtime_prompt_old_chatgpt.md')
    for path in candidates:
        if path.is_file():
            return path.read_text(encoding='utf-8'), path.name
    raise ValueError('Canonical runtime prompt is missing. Set RUNTIME_PROMPT_PATH to Runtime_prompt_old_chatgpt.md')

def _docx_text(path):
    from docx import Document
    document = Document(path)
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for index, table in enumerate(document.tables, 1):
        rows = []
        for row in table.rows:
            cells = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
            if any(cells):
                rows.append(' | '.join(cells))
        if rows:
            parts.append(f'Table {index}\n' + '\n'.join(rows))
    return '\n'.join(parts)

def _estimate_cost(settings, inp, cached, out):
    in_rate = float(getattr(settings, 'openai_input_usd_per_million', 0) or 0)
    cached_rate = float(getattr(settings, 'openai_cached_input_usd_per_million', 0) or 0)
    out_rate = float(getattr(settings, 'openai_output_usd_per_million', 0) or 0)
    return ((inp - cached) * in_rate + cached * cached_rate + out * out_rate) / 1e6

class Inspector:
    def __init__(self, settings, drive, client=None):
        self.settings = settings
        self.drive = drive
        self.diagnostics = Diagnostics(getattr(settings, 'diagnostics_dir', './logs'))
        if client is None:
            from openai import OpenAI
            client = OpenAI(api_key=settings.openai_api_key, max_retries=0, timeout=getattr(settings, 'openai_timeout_seconds', 900))
        self.client = client

    def inspect(self, job, batch, sources, context, rules):
        prompt_text, prompt_name = _runtime_prompt(self.settings)
        docs = Path(self.settings.governing_docs_dir)
        references = []
        doc_names = []
        for path in sorted(docs.iterdir()):
            if path.suffix.lower() in {'.md', '.txt'}:
                references.append(path.name + '\n' + path.read_text(encoding='utf-8'))
                doc_names.append(path.name)
            elif path.suffix.lower() == '.docx':
                references.append(path.name + '\n' + _docx_text(path))
                doc_names.append(path.name)
        if not references:
            raise ValueError('No readable governing documents')
        approved = context.get('approved_project_reference') if isinstance(context, dict) else None
        reference_id = ''
        reference_version = ''
        if isinstance(approved, dict):
            reference_id = approved.get('reference_id') or approved.get('source_file') or ''
            reference_version = str(approved.get('version') or '')
        images = []
        modes = []
        attempt = int(getattr(context, 'attempt', 0) or 0) if not isinstance(context, dict) else int(context.get('inspection_attempt', 0) or 0)
        start = time.monotonic()
        original_max_bytes = int(getattr(self.settings, 'image_original_max_bytes', ORIGINAL_MAX_BYTES) or ORIGINAL_MAX_BYTES)
        original_max_edge = int(getattr(self.settings, 'image_original_max_edge', ORIGINAL_MAX_EDGE) or ORIGINAL_MAX_EDGE)
        resize_max_edge = int(getattr(self.settings, 'image_resize_max_edge', RESIZE_MAX_EDGE) or RESIZE_MAX_EDGE)
        jpeg_quality = int(getattr(self.settings, 'image_jpeg_quality', JPEG_QUALITY) or JPEG_QUALITY)
        try:
            for src in sources:
                raw = self.drive.image(src)
                if hasattr(raw, 'read'):
                    raw = raw.read()
                data, mime, mode = prepare_model_image(
                    raw,
                    getattr(src, 'mime_type', '') or 'image/jpeg',
                    original_max_bytes=original_max_bytes,
                    original_max_edge=original_max_edge,
                    resize_max_edge=resize_max_edge,
                    jpeg_quality=jpeg_quality,
                )
                modes.append({'filename': src.name, 'mode': mode, 'mime': mime})
                images.append({'type': 'input_text', 'text': f'Source filename: {src.name}; Drive ID: {src.id}'})
                images.append({
                    'type': 'input_image',
                    'image_url': f'data:{mime};base64,{base64.b64encode(data).decode()}',
                    'detail': 'high',
                })
        except Exception as exc:
            setattr(exc, 'elapsed_seconds', round(time.monotonic() - start, 2))
            try:
                path = self.diagnostics.write(
                    job, batch, attempt, 'inspection_error', error_type=type(exc).__name__, error=str(exc),
                    elapsed_seconds=getattr(exc, 'elapsed_seconds', 0), success=False, image_preparation=modes,
                )
                setattr(exc, 'diagnostics_path', str(path))
            except Exception:
                pass
            raise
        instruction = (
            'The client runtime prompt below is the canonical inspection workflow. '
            'Return Lean Split v2 Asset Staging and Record Staging in the structured object. '
            'Do not update Google Sheets. '
            'A client-approved calibration example in the user payload is calibration for the named project and box, not a filename rule and not an answer-key extract. '
            'Inspect every authorized image. Do not copy an example when the images disagree, and do not close an item only because its folder has ended.\n\n'
            + prompt_text
            + '\n\n'
            + '\n\n'.join(references)
            + '\n\n'
            + OUTPUT_SCHEMA
        )
        prior = context.get('prior_open_item') if isinstance(context, dict) else None
        prior_text = 'None' if prior in (None, '', 'None') else str(prior)
        existing = []
        if isinstance(context, dict) and prior_text != 'None':
            existing = context.get('existing_open_record') or context.get('prior_open_item_context') or []
        payload = {
            'job_id': job,
            'batch_id': batch,
            'project': context.get('project') if isinstance(context, dict) else '',
            'run_context': context.get('run_context') if isinstance(context, dict) else [],
            'batch_queue_row': context.get('batch_queue_row') if isinstance(context, dict) else {},
            'manifest': context.get('manifest') if isinstance(context, dict) else [],
            'prior_open_item': prior_text,
            'prior_open_item_source': (context.get('prior_open_item_source') if isinstance(context, dict) else '') or 'Batch Queue Prior Open Item',
            'existing_open_record': existing,
            'prior_open_item_context': existing,
            'project_classification_rules': rules,
            'authorized_project_state': context.get('authorized_project_state') if isinstance(context, dict) else {},
            'approved_project_reference': approved,
            'folder_code': context.get('folder_code') if isinstance(context, dict) else '',
            'authorized_filenames': [s.name for s in sources],
        }
        start = time.monotonic()
        response_meta = {}
        attempt = int(getattr(context, 'attempt', 0) or 0) if not isinstance(context, dict) else int(context.get('inspection_attempt', 0) or 0)
        max_tokens = getattr(self.settings, 'openai_max_output_tokens', 0)
        effort = (getattr(self.settings, 'openai_reasoning_effort', '') or '').strip()
        prompt_ids = [prompt_name, *doc_names]
        if reference_id:
            prompt_ids.append(reference_id)
        self.diagnostics.write(
            job, batch, attempt, 'request_prepared',
            model=self.settings.openai_model,
            reasoning_effort=effort,
            max_output_tokens=max_tokens,
            source_image_count=len(sources),
            authorized_filenames=[s.name for s in sources],
            prior_open_item=prior_text,
            prior_open_item_source=payload['prior_open_item_source'],
            prompt_identifiers=prompt_ids,
            reference_id=reference_id,
            reference_version=reference_version,
            image_preparation=modes,
            prompt_character_count=len(instruction),
            rules_count=len(rules),
        )
        try:
            request = {
                'model': self.settings.openai_model,
                'input': [
                    {'role': 'system', 'content': instruction},
                    {'role': 'user', 'content': [{'type': 'input_text', 'text': json.dumps(payload)}] + images},
                ],
                'text_format': InspectionOutput,
            }
            if max_tokens:
                request['max_output_tokens'] = max_tokens
            if effort:
                request['reasoning'] = {'effort': effort}
            response = self.client.responses.parse(**request)
            usage = getattr(response, 'usage', None)
            details = getattr(usage, 'input_tokens_details', None)
            response_meta = {
                'response_id': getattr(response, 'id', ''),
                'status': getattr(response, 'status', ''),
                'incomplete_details': getattr(response, 'incomplete_details', None),
                'input_tokens': getattr(usage, 'input_tokens', 0) if usage else 0,
                'cached_input_tokens': getattr(details, 'cached_tokens', 0) if details else 0,
                'output_tokens': getattr(usage, 'output_tokens', 0) if usage else 0,
            }
            response_meta['estimated_cost_usd'] = _estimate_cost(
                self.settings,
                int(response_meta['input_tokens'] or 0),
                int(response_meta['cached_input_tokens'] or 0),
                int(response_meta['output_tokens'] or 0),
            )
            output = response.output_parsed
            returned_assets = [a.digital_asset for a in output.assets] if output else []
            returned_records = [r.proposed_record_id for r in output.records] if output else []
            self.diagnostics.write(
                job, batch, attempt, 'response_received', **response_meta,
                returned_asset_filenames=returned_assets,
                returned_record_ids=returned_records,
            )
            if output is None:
                raise ValueError('Missing structured output')
            expected = [s.name for s in sources]
            actual = [a.digital_asset for a in output.assets]
            if actual != expected:
                missing = sorted(set(expected) - set(actual))
                extra = sorted(set(actual) - set(expected))
                self.diagnostics.write(
                    job, batch, attempt, 'inspection_validation_failed',
                    reason='filename_integrity', expected_count=len(expected),
                    returned_count=len(actual), missing=missing, extra=extra,
                    expected_order=expected, returned_order=actual,
                )
                raise ValueError(
                    f'Inspection returned incomplete/non-ordered staging: expected {len(expected)} assets, got {len(actual)}; missing={missing[:5]}; extra={extra[:5]}'
                )
            if len(output.records) == 0:
                self.diagnostics.write(job, batch, attempt, 'inspection_validation_failed', reason='no_record_rows')
                raise ValueError('Inspection returned no Record Staging rows')
            linked = {a.assigned_record_id for a in output.assets}
            proposed = {r.proposed_record_id for r in output.records}
            if not all(a.assigned_record_id for a in output.assets) or not linked.issubset(proposed):
                self.diagnostics.write(
                    job, batch, attempt, 'inspection_validation_failed', reason='unlinked_assets',
                    linked_record_ids=sorted(linked), returned_record_ids=sorted(proposed),
                )
                raise ValueError('Every inspected asset must link to a returned Record Staging row')
            elapsed = round(time.monotonic() - start, 2)
            self.diagnostics.write(
                job, batch, attempt, 'inspection_validated',
                elapsed_seconds=elapsed,
                success=True,
                image_preparation=modes,
                estimated_cost_usd=response_meta['estimated_cost_usd'],
                response_id=response_meta['response_id'],
            )
            usage_out = {
                'input_tokens': response_meta['input_tokens'],
                'cached_input_tokens': response_meta['cached_input_tokens'] or 0,
                'output_tokens': response_meta['output_tokens'],
                'elapsed_seconds': elapsed,
                'response_id': response_meta['response_id'],
                'estimated_cost_usd': response_meta['estimated_cost_usd'],
            }
            return output, usage_out
        except Exception as e:
            setattr(e, 'elapsed_seconds', round(time.monotonic() - start, 2))
            if response_meta:
                setattr(e, 'inspection_usage', {
                    'input_tokens': response_meta.get('input_tokens') or 0,
                    'cached_input_tokens': response_meta.get('cached_input_tokens') or 0,
                    'output_tokens': response_meta.get('output_tokens') or 0,
                    'response_id': response_meta.get('response_id') or '',
                    'elapsed_seconds': getattr(e, 'elapsed_seconds', 0),
                    'estimated_cost_usd': response_meta.get('estimated_cost_usd') or 0,
                })
            try:
                path = self.diagnostics.write(
                    job, batch, attempt, 'inspection_error',
                    error_type=type(e).__name__,
                    error=str(e),
                    elapsed_seconds=getattr(e, 'elapsed_seconds', 0),
                    success=False,
                    image_preparation=modes,
                    prior_open_item_source=payload['prior_open_item_source'],
                    prompt_identifiers=prompt_ids,
                    estimated_cost_usd=(response_meta or {}).get('estimated_cost_usd', 0),
                    response_id=(response_meta or {}).get('response_id', ''),
                )
                setattr(e, 'diagnostics_path', str(path))
            except Exception:
                pass
            credit = as_credit_error(e)
            if credit is not e:
                raise credit from e
            raise
