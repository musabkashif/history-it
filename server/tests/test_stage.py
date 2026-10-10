import json
import tempfile
import unittest
from pathlib import Path

from client_reference.commit_engine import CommitError
from app.models import ASSET_HEADERS, RECORD_HEADERS, AssetRow, InspectionOutput, RecordRow
from app.orchestrator import Orchestrator, unit_size


class MemorySheets:
    def __init__(self):
        self.store = {}

    def rows(self, sid, sheet):
        return [dict(row) for row in self.store.get((sid, sheet), [])]

    def write(self, sid, sheet, row, key=None):
        rows = self.store.setdefault((sid, sheet), [])
        if isinstance(key, list):
            matches = [existing for existing in rows if all(str(existing.get(name)) == str(value) for name, value in key)]
        elif key:
            matches = [existing for existing in rows if str(existing.get(key[0])) == str(key[1])]
        else:
            matches = []
        if matches:
            matches[0].update({name: str(value) for name, value in row.items()})
            return 1
        rows.append({name: str(value) for name, value in row.items()})
        return len(rows)

    def clear(self, sid, sheet, row_number, key):
        rows = self.store.get((sid, sheet), [])
        if isinstance(key, list):
            self.store[(sid, sheet)] = [
                row for row in rows if not all(str(row.get(name)) == str(value) for name, value in key)
            ]
        else:
            self.store[(sid, sheet)] = [row for row in rows if str(row.get(key[0])) != str(key[1])]

    def field(self, sid, name):
        return 'None'

    def ensure_tab(self, sid, name, headers):
        self.store.setdefault((sid, name), [])


def filled(headers, **values):
    row = {name: '' for name in headers}
    row.update(values)
    return row


def batch():
    names = ['TriDelta_Year2_Chapters_Box19_F01_001.jpg', 'TriDelta_Year2_Chapters_Box19_F01_002.jpg']
    manifest = [{'id': name, 'name': name, 'folder_id': 'folder', 'folder_name': '1vNkAYMTmohyqR4ebP4DJ_BIjN_R-AhJ5', 'path': '/' + name, 'mime_type': 'image/jpeg'} for name in names]
    return {
        'Project': 'Tri Delta Chapters',
        'Box': 'Chapters Box 19',
        'Source Group / Folder': '1vNkAYMTmohyqR4ebP4DJ_BIjN_R-AhJ5',
        'Job ID': 'HIT-f0ffaf9f935c905b',
        'Batch ID': 'HIT-f0ffaf9f935c905b-B0001',
        'Planned Asset Count': '2',
        'Notes': json.dumps({'manifest': manifest}),
    }


class StageControlTests(unittest.TestCase):
    def test_model_bookkeeping_is_replaced_before_write(self):
        sheets = MemorySheets()
        orch = Orchestrator(sheets, None, None, None)
        b = batch()
        assets = []
        for name in ['TriDelta_Year2_Chapters_Box19_F01_001.jpg', 'TriDelta_Year2_Chapters_Box19_F01_002.jpg']:
            assets.append(AssetRow.model_validate(filled(
                ASSET_HEADERS,
                **{
                    'Project': 'wrong',
                    'Job ID': '',
                    'Batch ID': 'Pending Authorization',
                    'Digital Asset': name,
                    'Assigned Record ID': 'TriDelta_Ch19_F01_R001',
                    'Source Group / Folder': 'F01 - Not Started',
                    'QC Status': 'Crop required',
                    'Inspection Complete': 'Complete',
                    'Commit Status': '',
                },
            )))
        record = RecordRow.model_validate(filled(
            RECORD_HEADERS,
            **{
                'Project': '',
                'Job ID': 'chat',
                'Batch ID': '',
                'Proposed Record ID': 'TriDelta_Ch19_F01_R001',
                'Source Group / Folder': 'F01 - Not Started',
                'Record Kind': 'Normal Odyssey',
                'Record Operation': 'Create New',
                'Metadata Readiness': 'Ready',
                'Template': 'Document',
                'Inspection Complete': 'done',
                'Commit Status': 'Pending',
            },
        ))
        orch._stage('sheet', b, InspectionOutput(assets=assets, records=[record]))
        written_assets = sheets.rows('sheet', 'Asset Staging')
        written_records = sheets.rows('sheet', 'Record Staging')
        self.assertEqual([row['Digital Asset'] for row in written_assets], [
            'TriDelta_Year2_Chapters_Box19_F01_001.jpg',
            'TriDelta_Year2_Chapters_Box19_F01_002.jpg',
        ])
        self.assertEqual(written_assets[0]['QC Status'], 'Crop required')
        self.assertEqual(written_assets[0]['Source Group / Folder'], 'F01')
        self.assertEqual(written_assets[0]['Assigned Record ID'], 'TDC-B19-F01-001')
        self.assertEqual(written_records[0]['Record Kind'], 'Normal Odyssey')
        self.assertEqual(written_records[0]['Template'], 'Document')
        self.assertEqual(written_records[0]['Proposed Record ID'], 'TDC-B19-F01-001')
        self.assertEqual(sheets.rows('sheet', 'Batch Queue')[0]['Source Group / Folder'], 'F01')
        for row in written_assets + written_records:
            self.assertEqual(row['Project'], 'Tri Delta Chapters')
            self.assertEqual(row['Job ID'], 'HIT-f0ffaf9f935c905b')
            self.assertEqual(row['Batch ID'], 'HIT-f0ffaf9f935c905b-B0001')
            self.assertEqual(row['Inspection Complete'], 'Yes')
            self.assertEqual(row['Commit Status'], 'Pending Commit')

    def test_blank_record_link_is_not_written(self):
        sheets = MemorySheets()
        orch = Orchestrator(sheets, None, None, None)
        b = batch()
        assets = [
            AssetRow.model_validate(filled(ASSET_HEADERS, **{'Digital Asset': name, 'Assigned Record ID': ''}))
            for name in ['TriDelta_Year2_Chapters_Box19_F01_001.jpg', 'TriDelta_Year2_Chapters_Box19_F01_002.jpg']
        ]
        with self.assertRaises(CommitError):
            orch._stage('sheet', b, InspectionOutput(assets=assets, records=[]))
        self.assertEqual(sheets.rows('sheet', 'Asset Staging'), [])
        self.assertEqual(sheets.rows('sheet', 'Record Staging'), [])

    def test_commit_failure_with_unlinked_assets_is_inspected_again(self):
        sheets = MemorySheets()
        b = batch()
        b['Batch Status'] = 'Commit Failed - Retryable'
        b['Inspection Attempt'] = '2'
        b['Error / Retry Note'] = 'Asset TriDelta_Year2_Chapters_Box19_F01_001.jpg refers to unstaged record '
        sheets.store[('index', 'HistoryIT Jobs')] = [{'Job ID': b['Job ID'], 'Spreadsheet ID': 'register'}]
        sheets.store[('register', 'Batch Queue')] = [b]
        sheets.store[('register', 'Asset Staging')] = [
            {'Batch ID': b['Batch ID'], 'Digital Asset': name, 'Assigned Record ID': '', 'Inspection Complete': 'Yes', 'Commit Status': 'Pending Commit'}
            for name in ['TriDelta_Year2_Chapters_Box19_F01_001.jpg', 'TriDelta_Year2_Chapters_Box19_F01_002.jpg']
        ]
        calls = []

        class Inspector:
            def inspect(self, job, batch_id, sources, context, rules):
                calls.append(batch_id)
                raise RuntimeError('reinspected')

        class Settings:
            job_index_spreadsheet_id = 'index'
            job_spending_limit_usd = 100
            max_retries = 3
            openai_model = 'gpt-test'
            openai_input_usd_per_million = 1
            openai_cached_input_usd_per_million = 0
            openai_output_usd_per_million = 1

        Orchestrator(sheets, None, Inspector(), Settings()).run(b['Job ID'])
        self.assertEqual(calls, [b['Batch ID']])
        row = sheets.rows('register', 'Batch Queue')[0]
        self.assertEqual(row['Batch Status'], 'Inspection Failed - Retryable')
        self.assertEqual(row['Error / Retry Note'], 'reinspected')

    def test_filename_order_is_still_rejected(self):
        sheets = MemorySheets()
        orch = Orchestrator(sheets, None, None, None)
        asset = AssetRow.model_validate(filled(ASSET_HEADERS, **{'Digital Asset': 'other.jpg'}))
        with self.assertRaisesRegex(ValueError, 'Staging filename/order mismatch'):
            orch._stage('sheet', batch(), InspectionOutput(assets=[asset], records=[]))
        self.assertEqual(sheets.rows('sheet', 'Asset Staging'), [])

    def test_metadata_reference_gets_register_id_and_no_task(self):
        sheets = MemorySheets()
        names = ['TriDelta_Year2_Chapters_Box19_F01_001.jpg', 'TriDelta_Year2_Chapters_Box19_F01_002.jpg']
        assets = [
            AssetRow.model_validate(filled(ASSET_HEADERS, **{
                'Digital Asset': name,
                'Assigned Record ID': 'cover' if name.endswith('001.jpg') else 'sheet',
                'QC Status': 'Pass',
                'Production Queue Eligible': 'Yes',
                'Expected Output Asset Count': '1',
            }))
            for name in names
        ]
        records = [
            RecordRow.model_validate(filled(RECORD_HEADERS, **{
                'Proposed Record ID': 'cover',
                'Record Kind': 'Metadata Only',
                'Record Operation': 'Create New',
                'Template': 'Document',
                'Metadata Readiness': 'Ready',
            })),
            RecordRow.model_validate(filled(RECORD_HEADERS, **{
                'Proposed Record ID': 'sheet',
                'Record Kind': 'Metadata Only',
                'Record Operation': 'Create New',
                'Template': 'Document',
            })),
        ]
        Orchestrator(sheets, None, None, None)._stage('sheet', batch(), InspectionOutput(assets=assets, records=records))
        written = sheets.rows('sheet', 'Record Staging')
        self.assertEqual([row['Proposed Record ID'] for row in written], ['TDC-B19-F01-REF-001', 'TDC-B19-F01-REF-002'])
        assets_written = sheets.rows('sheet', 'Asset Staging')
        self.assertEqual(assets_written[1]['Assigned Record ID'], 'TDC-B19-F01-REF-002')
        self.assertEqual(assets_written[1]['QC Status'], 'Reference Only')
        self.assertEqual(assets_written[1]['Action Needed'], 'N/A')
        self.assertEqual(assets_written[1]['Production Queue Eligible'], 'No')
        self.assertEqual(assets_written[1]['Expected Output Asset Count'], '0')
        self.assertEqual(written[1]['Template'], 'Metadata Only')
        self.assertEqual(written[1]['Metadata Readiness'], 'Metadata Reference')
        self.assertEqual(written[1]['Production Task Count'], '0')

    def test_no_adjustment_does_not_create_a_task(self):
        sheets = MemorySheets()
        names = ['TriDelta_Year2_Chapters_Box19_F01_001.jpg', 'TriDelta_Year2_Chapters_Box19_F01_002.jpg']
        assets = [
            AssetRow.model_validate(filled(ASSET_HEADERS, **{
                'Digital Asset': name,
                'Assigned Record ID': 'item',
                'QC Status': 'Pass',
                'Action Needed': 'No Adjustment',
                'Expected Output Asset Count': '1',
                'Production Queue Eligible': 'Yes',
                'Production Task Scope': 'crop',
                'Production Task Note': 'queued by model',
            }))
            for name in names
        ]
        record = RecordRow.model_validate(filled(RECORD_HEADERS, **{
            'Proposed Record ID': 'item',
            'Record Kind': 'Normal Odyssey',
            'Record Operation': 'Create New',
            'Metadata Readiness': 'Ready',
            'Production Task Count': '3',
        }))
        Orchestrator(sheets, None, None, None)._stage('sheet', batch(), InspectionOutput(assets=assets, records=[record]))
        written_assets = sheets.rows('sheet', 'Asset Staging')
        self.assertEqual({row['Production Queue Eligible'] for row in written_assets}, {'No'})
        self.assertEqual({row['Expected Output Asset Count'] for row in written_assets}, {'1'})
        self.assertEqual({row['Production Task Scope'] for row in written_assets}, {''})
        self.assertEqual(sheets.rows('sheet', 'Record Staging')[0]['Production Task Count'], '0')

    def test_crop_action_is_queued_and_prior_tasks_carry_forward(self):
        sheets = MemorySheets()
        sheets.store[('sheet', 'Itemization')] = [{
            'Item ID': 'TDC-B19-F01-001',
            'Digital Assets': 'TriDelta_Year2_Chapters_Box19_F01_001.jpg',
            'Production Task Count': '2',
        }]
        names = ['TriDelta_Year2_Chapters_Box19_F01_004.jpg', 'TriDelta_Year2_Chapters_Box19_F01_005.jpg']
        queued = batch()
        queued['Notes'] = json.dumps({'manifest': [{
            'id': name, 'name': name, 'folder_id': 'folder',
            'folder_name': 'F01', 'path': '/' + name, 'mime_type': 'image/jpeg',
        } for name in names]})
        assets = [
            AssetRow.model_validate(filled(ASSET_HEADERS, **{
                'Digital Asset': names[0],
                'Assigned Record ID': 'TDC-B19-F01-001',
                'Action Needed': 'No Adjustment',
                'Production Queue Eligible': 'Yes',
                'Expected Output Asset Count': '1',
            })),
            AssetRow.model_validate(filled(ASSET_HEADERS, **{
                'Digital Asset': names[1],
                'Assigned Record ID': 'TDC-B19-F01-001',
                'Action Needed': 'Crop / Recrop',
                'Production Queue Eligible': 'No',
                'Expected Output Asset Count': '1',
            })),
        ]
        record = RecordRow.model_validate(filled(RECORD_HEADERS, **{
            'Proposed Record ID': 'TDC-B19-F01-001',
            'Record Kind': 'Normal Odyssey',
            'Record Operation': 'Update Existing Open Record',
            'Metadata Readiness': 'Blocked - Cropping/QC',
            'Production Task Count': '0',
        }))
        Orchestrator(sheets, None, None, None)._stage('sheet', queued, InspectionOutput(assets=assets, records=[record]), 'TDC-B19-F01-001')
        written = sheets.rows('sheet', 'Asset Staging')
        self.assertEqual(written[0]['Production Queue Eligible'], 'No')
        self.assertEqual(written[1]['Production Queue Eligible'], 'Yes')
        self.assertEqual(sheets.rows('sheet', 'Record Staging')[0]['Production Task Count'], '3')

    def test_folder_within_client_batch_range_stays_whole(self):
        self.assertEqual(unit_size(6, 21), 6)
        self.assertEqual(unit_size(2, 25), 2)
        self.assertEqual(unit_size(6, 100), 6)
        self.assertEqual(unit_size(15, 100), 15)
        self.assertEqual(unit_size(None, 40), 10)

    def test_create_uses_the_client_batch_size(self):
        from app.models import JobCreate, Source

        class Drive:
            def __init__(self, count):
                self.count = count

            def enumerate(self, root):
                return [
                    Source(id=str(i), name=f'TriDelta_Year2_Chapters_Box19_F01_{i:03d}.jpg', folder_id='folder', folder_name='F01', path=f'/{i}', mime_type='image/jpeg')
                    for i in range(1, self.count + 1)
                ]

        class Settings:
            job_index_spreadsheet_id = 'index'
            batch_size = 2

        def created(count, requested):
            sheets = MemorySheets()
            Orchestrator(sheets, Drive(count), None, Settings()).create(JobCreate(
                project='Tri Delta Chapters', source_folder_id='folder', spreadsheet_id='register', batch_size=requested, box='Chapters Box 19',
            ))
            return [row['Planned Asset Count'] for row in sheets.rows('register', 'Batch Queue')]

        self.assertEqual(created(21, 6), ['6', '6', '6', '3'])
        self.assertEqual(created(30, 6), ['6', '6', '6', '6', '6'])

    def test_open_reference_keeps_its_id_and_assets(self):
        sheets = MemorySheets()
        prior = 'TDC-B19-F01-REF-002'
        sheets.store[('sheet', 'Metadata References')] = [{
            'Reference ID': prior,
            'Digital Assets': 'TriDelta_Year2_Chapters_Box19_F01_002.jpg',
        }]
        name = 'TriDelta_Year2_Chapters_Box19_F01_003.jpg'
        queued = batch()
        queued['Notes'] = json.dumps({'manifest': [{
            'id': name, 'name': name, 'folder_id': 'folder',
            'folder_name': '1vNkAYMTmohyqR4ebP4DJ_BIjN_R-AhJ5', 'path': '/' + name, 'mime_type': 'image/jpeg',
        }]})
        queued['Planned Asset Count'] = '1'
        asset = AssetRow.model_validate(filled(ASSET_HEADERS, **{'Digital Asset': name, 'Assigned Record ID': 'local'}))
        record = RecordRow.model_validate(filled(RECORD_HEADERS, **{
            'Proposed Record ID': 'local',
            'Record Kind': 'Metadata Only',
            'Record Operation': 'Update Existing Open Record',
        }))
        Orchestrator(sheets, None, None, None)._stage('sheet', queued, InspectionOutput(assets=[asset], records=[record]), prior)
        written = sheets.rows('sheet', 'Record Staging')[0]
        self.assertEqual(written['Proposed Record ID'], prior)
        self.assertEqual(written['Record Operation'], 'Update Existing Open Record')
        self.assertEqual(written['Full Record Assets After Commit'], 'TriDelta_Year2_Chapters_Box19_F01_002.jpg; ' + name)
        self.assertEqual(sheets.rows('sheet', 'Asset Staging')[0]['Assigned Record ID'], prior)

    def test_replaced_record_id_does_not_remain(self):
        sheets = MemorySheets()
        sheets.store[('sheet', 'Record Staging')] = [{
            'Batch ID': 'HIT-f0ffaf9f935c905b-B0001',
            'Proposed Record ID': 'TriDelta_Ch19_F01_R001',
            '_row': 8,
        }]
        names = ['TriDelta_Year2_Chapters_Box19_F01_001.jpg', 'TriDelta_Year2_Chapters_Box19_F01_002.jpg']
        assets = [AssetRow.model_validate(filled(ASSET_HEADERS, **{'Digital Asset': name, 'Assigned Record ID': 'item'})) for name in names]
        record = RecordRow.model_validate(filled(RECORD_HEADERS, **{
            'Proposed Record ID': 'item',
            'Record Kind': 'Normal Odyssey',
            'Record Operation': 'Create New',
            'Metadata Readiness': 'Ready',
        }))
        Orchestrator(sheets, None, None, None)._stage('sheet', batch(), InspectionOutput(assets=assets, records=[record]))
        ids = [row['Proposed Record ID'] for row in sheets.rows('sheet', 'Record Staging')]
        self.assertEqual(ids, ['TDC-B19-F01-001'])

    def test_governing_tables_are_sent_to_inspection(self):
        from docx import Document
        from app.inspection import Inspector

        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder, 'rules.docx')
            document = Document()
            document.add_paragraph('Folder covers are Metadata Only.')
            table = document.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = 'Action'
            table.rows[0].cells[1].text = 'Archival Border Crop'
            document.save(path)

            class Settings:
                governing_docs_dir = folder
                openai_api_key = 'test'
                openai_model = 'gpt-test'

            class Drive:
                def image(self, source):
                    from io import BytesIO
                    from PIL import Image
                    out = BytesIO()
                    Image.new('RGB', (3000, 1200), 'white').save(out, format='JPEG')
                    return out.getvalue()

            captured = {}

            class Client:
                def __init__(self):
                    self.responses = self

                def parse(self, **kwargs):
                    captured.update(kwargs)
                    raise RuntimeError('captured')

            class SourceFile:
                name = 'a.jpg'
                id = 'file-1'
                mime_type = 'image/jpeg'

            with self.assertRaisesRegex(RuntimeError, 'captured'):
                Inspector(Settings(), Drive(), client=Client()).inspect('HIT-1', 'B1', [SourceFile()], {'folder_code': 'F01', 'prior_open_item': 'None'}, [])
            instruction = captured['input'][0]['content']
            self.assertIn('INSPECTION PHASE ONLY', instruction)
            self.assertIn('Archival Border Crop', instruction)
            self.assertIn('technical batch boundary does not force', instruction)
            self.assertNotIn('blank back stays on that same Metadata Only reference', instruction)
            self.assertNotIn('F01-Report', instruction)
            self.assertNotIn('Frozen Benchmark', instruction)
            self.assertIn('Folder covers are Metadata Only.', instruction)
            image = captured['input'][1]['content'][2]
            self.assertTrue(image['image_url'].startswith('data:image/jpeg;base64,'))
            self.assertEqual(image['detail'], 'high')
            payload = json.loads(captured['input'][1]['content'][0]['text'])
            self.assertIsNone(payload['approved_project_reference'])
            self.assertNotIn('Frozen Benchmark', captured['input'][1]['content'][0]['text'])

    def test_original_image_is_not_crushed_to_a_lossy_preview(self):
        from io import BytesIO
        from PIL import Image
        from app.images import prepare_model_image
        source = BytesIO()
        Image.new('RGB', (1600, 400), 'red').save(source, format='JPEG', quality=95)
        original = source.getvalue()
        kept, mime, mode = prepare_model_image(original, 'image/jpeg')
        self.assertEqual(mode, 'original')
        self.assertEqual(mime, 'image/jpeg')
        self.assertEqual(kept, original)
        with Image.open(BytesIO(kept)) as image:
            self.assertEqual(image.size, (1600, 400))
        large = BytesIO()
        Image.new('RGB', (6000, 3000), 'blue').save(large, format='JPEG', quality=95)
        reduced, reduced_mime, reduced_mode = prepare_model_image(large.getvalue(), 'image/jpeg')
        self.assertEqual(reduced_mode, 'resized-high-quality')
        self.assertEqual(reduced_mime, 'image/jpeg')
        self.assertLess(len(reduced), len(large.getvalue()))
        with Image.open(BytesIO(reduced)) as image:
            self.assertEqual(image.size, (2048, 1024))
            self.assertEqual(round(image.size[0] / image.size[1], 2), round(6000 / 3000, 2))

    def test_dropped_connection_is_retried_and_a_client_error_is_not(self):
        from threading import Lock
        from app.adapters import _execute

        calls = []

        class Download:
            def execute(self):
                calls.append('download')
                if len(calls) < 3:
                    raise ConnectionResetError(10054, 'An existing connection was forcibly closed by the remote host')
                return b'jpeg'

        self.assertEqual(_execute(Lock(), lambda: Download(), sleep=lambda _seconds: None), b'jpeg')
        self.assertEqual(len(calls), 3)

        class Denied(Exception):
            def __init__(self):
                self.resp = type('Response', (), {'status': 400})()

        denied = []

        class Rejected:
            def execute(self):
                denied.append(1)
                raise Denied()

        with self.assertRaises(Denied):
            _execute(Lock(), lambda: Rejected(), sleep=lambda _seconds: None)
        self.assertEqual(denied, [1])


if __name__ == '__main__':
    unittest.main()
