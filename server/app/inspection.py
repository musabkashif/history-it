from __future__ import annotations
import base64, json, time
from pathlib import Path
from app.adapters import IncompleteDriveDownload
from app.credits import as_credit_error
from app.diagnostics import Diagnostics, exception_chain
from app.images import JPEG_QUALITY, ORIGINAL_MAX_BYTES, ORIGINAL_MAX_EDGE, RESIZE_MAX_EDGE, prepare_model_image
from app.decisions import decision_instructions
from app.models import InspectionOutput

OUTPUT_SCHEMA = """Structured output schema only. This text does not decide record boundaries, record types, templates, or metadata readiness. Those judgments come from the client runtime prompt, the governing documents, the authorized images, and any client-approved calibration example in the user payload.

Return Lean Split v2 Asset Staging and Record Staging as the structured object. Do not update a spreadsheet, choose a batch, or calculate cumulative totals. All values must be strings. Inspect every authorized image. Preserve every authorized filename exactly once, in manifest order. Each Assigned Record ID must equal a Proposed Record ID in this response. Record Kind is Normal Odyssey or Metadata Only. Record Operation is Create New or Update Existing Open Record. Copy context.prior_open_item from the payload when continuing that record; do not replace a None value and do not create a second record for it. Copy context.folder_code exactly.

When Boundary State After Record starts with Open, or Metadata Readiness is Review, write the visual evidence for that decision in Review Note. Leave Review Note blank only when the record is neither Review nor Open. A folder boundary does not by itself close an item."""

CONNECTION_ATTEMPTS = 4
# Drive media only. Separate from OpenAI connection retries and never unbounded.
DRIVE_DOWNLOAD_ATTEMPTS = 3
DRIVE_DOWNLOAD_BACKOFF_SECONDS = 1
# Provider HTTP 500 binary splits. Depth is capped; size 1 cannot split further.
PROVIDER_SPLIT_MAX_DEPTH = 6
PROVIDER_SPLIT_MIN_IMAGES = 2

_AMBIGUOUS_TIMEOUT_NAMES = {
    'APITimeoutError', 'ReadTimeout', 'WriteTimeout', 'PoolTimeout', 'TimeoutError',
}

_DRIVE_DOWNLOAD_ERROR_NAMES = {
    'IncompleteRead', 'IncompleteDriveDownload', 'SSLEOFError', 'SSLError',
    'RemoteDisconnected', 'ConnectionResetError', 'ConnectionAbortedError',
    'BrokenPipeError', 'ProtocolError', 'TimeoutError',
}


def _walk_exceptions(exc):
    seen = set()
    current = exc
    while current is not None and id(current) not in seen and len(seen) < 8:
        seen.add(id(current))
        yield current
        nxt = current.__cause__
        if nxt is None and not getattr(current, '__suppress_context__', False):
            nxt = current.__context__
        current = nxt


def _is_ambiguous_timeout(exc):
    """True when the provider may already have accepted the request."""
    for current in _walk_exceptions(exc):
        names = {cls.__name__ for cls in type(current).__mro__}
        if names & {'ConnectTimeout', 'ConnectError'}:
            continue
        if names & _AMBIGUOUS_TIMEOUT_NAMES:
            return True
    return False


def _is_connection_error(exc):
    if _is_ambiguous_timeout(exc):
        return False
    for current in _walk_exceptions(exc):
        names = {cls.__name__ for cls in type(current).__mro__}
        if 'APITimeoutError' not in names and names & {
            'APIConnectionError', 'ConnectError', 'ReadError', 'WriteError',
            'RemoteProtocolError', 'ConnectionError', 'ConnectionResetError',
            'gaierror', 'TimeoutError',
        }:
            return True
    return False


def _is_drive_download_error(exc):
    """True for idempotent Drive media transport failures. Never OpenAI timeouts."""
    if isinstance(exc, IncompleteDriveDownload):
        return True
    for current in _walk_exceptions(exc):
        names = {cls.__name__ for cls in type(current).__mro__}
        if names & {'APITimeoutError', 'ReadTimeout', 'WriteTimeout', 'PoolTimeout'}:
            return False
        if names & _DRIVE_DOWNLOAD_ERROR_NAMES:
            return True
        text = f'{type(current).__name__} {current}'.casefold()
        if 'incomplete read' in text or 'ssleof' in text or 'violation of protocol' in text:
            return True
    return False


def provider_server_error(exc) -> bool:
    """True when the provider failed the request and a smaller request may succeed."""
    seen = set()
    current = exc
    while current is not None and id(current) not in seen and len(seen) < 8:
        seen.add(id(current))
        status = getattr(current, 'status_code', None)
        if status in (500, 502, 503):
            return True
        names = {cls.__name__ for cls in type(current).__mro__}
        if 'InternalServerError' in names:
            return True
        nxt = current.__cause__
        if nxt is None and not current.__suppress_context__:
            nxt = current.__context__
        current = nxt
    return False


def transient_inspection_failure(note: str) -> bool:
    """Network drops and provider HTTP 500s can be retried. They are not content failures."""
    text = (note or '').lower()
    return (
        'getaddrinfo failed' in text
        or 'connection error' in text
        or 'error code: 500' in text
        or 'server had an error processing your request' in text
        or 'provider split depth limit' in text
        or 'provider_split_max_depth' in text
        or 'part_below_minimum_split_size' in text
        or text.strip() == 'inspection retry limit reached'
    )


def inspection_attempt_blocked(attempt: int, max_retries: int, note: str) -> bool:
    """Content failures stop at the cap. A transient failure stays retryable."""
    if attempt <= max_retries:
        return False
    return not transient_inspection_failure(note)


def _still_open(record) -> bool:
    boundary = (record.boundary_state or '').casefold()
    readiness = (record.metadata_readiness or '').casefold()
    return boundary.startswith('open') or readiness == 'review'


def continuation_context(context, output):
    """Hand the open record at the end of one piece to the next piece of the same batch."""
    copied = dict(context or {})
    if not output.assets:
        return copied
    last_id = output.assets[-1].assigned_record_id
    match = next((record for record in output.records if record.proposed_record_id == last_id and _still_open(record)), None)
    if match is None:
        return copied
    copied['prior_open_item'] = last_id
    copied['prior_open_item_source'] = 'Previous images in this batch'
    detail = [match.model_dump(by_alias=True)]
    copied['existing_open_record'] = detail
    copied['prior_open_item_context'] = detail
    return copied


def _merge_record_operation(previous, incoming):
    """Keep Create New when a later split part Updates the same in-batch record."""
    previous_op = (previous.record_operation or '').strip()
    incoming_op = (incoming.record_operation or '').strip()
    if previous_op == 'Create New' and incoming_op == 'Update Existing Open Record':
        incoming.record_operation = 'Create New'
    return incoming


def merge_inspections(parts, continuation_ids=()):
    """Keep manifest order. A continued record stays one row. Other id clashes stay distinct.

    Within-batch continuation and repeated Update Existing rows for the same local
    id collapse to one record. Create-then-Update across split parts stays Create
    New so identity assignment does not require a committed Sheets row yet.
    """
    continued = {item for item in continuation_ids if item and item != 'None'}
    assets = []
    records = []
    seen = set()
    for part_index, output in enumerate(parts):
        remap = {}
        for record in output.records:
            local = record.proposed_record_id
            previous = next((item for item in records if item.proposed_record_id == local), None)
            same_prior_update = (
                previous is not None
                and (previous.record_operation or '').strip() == 'Update Existing Open Record'
                and (record.record_operation or '').strip() == 'Update Existing Open Record'
            )
            if local in seen and (local in continued or same_prior_update):
                if previous is not None:
                    record = _merge_record_operation(previous, record)
                records = [item for item in records if item.proposed_record_id != local]
                records.append(record)
                continue
            if local in seen:
                number = 1
                renamed = f'{local}~{number}'
                while renamed in seen:
                    number += 1
                    renamed = f'{local}~{number}'
                remap[local] = renamed
                record.proposed_record_id = renamed
                local = renamed
            seen.add(local)
            records.append(record)
        for asset in output.assets:
            if asset.assigned_record_id in remap:
                asset.assigned_record_id = remap[asset.assigned_record_id]
            assets.append(asset)
    return InspectionOutput(assets=assets, records=records)


def merge_usage(parts):
    combined = {
        'input_tokens': 0,
        'cached_input_tokens': 0,
        'cache_write_30m_tokens': 0,
        'cache_write_12h_tokens': 0,
        'output_tokens': 0,
        'elapsed_seconds': 0,
        'estimated_cost_usd': 0,
        'response_id': '',
        'responses': list(parts),
        'usage_recorded': bool(parts) and all(part.get('usage_recorded') for part in parts),
    }
    ids = []
    for part in parts:
        for key in ('input_tokens', 'cached_input_tokens', 'cache_write_30m_tokens', 'cache_write_12h_tokens', 'output_tokens'):
            combined[key] += int(part.get(key) or 0)
        combined['elapsed_seconds'] += float(part.get('elapsed_seconds') or 0)
        combined['estimated_cost_usd'] += float(part.get('estimated_cost_usd') or 0)
        if part.get('response_id'):
            ids.append(str(part['response_id']))
    combined['elapsed_seconds'] = round(combined['elapsed_seconds'], 2)
    combined['response_id'] = ','.join(ids)
    return combined


def _next_part(context) -> int:
    if not isinstance(context, dict):
        return 1
    holder = context.get('_usage_parts')
    if not isinstance(holder, dict):
        holder = {'n': 0}
        context['_usage_parts'] = holder
    holder['n'] += 1
    return holder['n']


class ProviderSplitExhausted(RuntimeError):
    """Provider splits reached the configured depth or minimum part size."""


def _split_tree(context) -> list:
    if not isinstance(context, dict):
        return []
    tree = context.setdefault('_provider_split_tree', [])
    if not isinstance(tree, list):
        context['_provider_split_tree'] = []
        return context['_provider_split_tree']
    return tree


def _split_node_id(context, depth: int, parent_id: str | None, branch: str | None) -> str:
    if parent_id and branch:
        return f'{parent_id}.{branch}'
    if isinstance(context, dict):
        holder = context.setdefault('_split_node_seq', {'n': 0})
        holder['n'] = int(holder.get('n') or 0) + 1
        return f'split-{holder["n"]}'
    return f'split-d{depth}'


def _provider_http_status(exc) -> int | None:
    for current in _walk_exceptions(exc):
        status = getattr(current, 'status_code', None)
        if status is not None:
            try:
                return int(status)
            except (TypeError, ValueError):
                return None
    return None


def _usage_billing_fields(usage: dict | None) -> dict:
    """Record only known usage. Do not invent tokens or billing."""
    usage = usage or {}
    response_id = str(usage.get('response_id') or '').strip()
    token_keys = (
        'input_tokens', 'cached_input_tokens', 'cache_write_tokens',
        'cache_write_30m_tokens', 'cache_write_12h_tokens', 'output_tokens',
    )
    tokens = {}
    for key in token_keys:
        if usage.get(key) not in (None, ''):
            tokens[key] = usage.get(key)
    if response_id:
        billing_status = 'billed_response'
        billed = True
    elif tokens:
        billing_status = 'unknown'
        billed = None
    else:
        billing_status = 'no_response'
        billed = False
    return {
        'response_id': response_id or None,
        'billed': billed,
        'billing_status': billing_status,
        'token_usage': tokens or None,
    }


def _emit_split_event(on_event, on_split, context, event: dict) -> None:
    tree = _split_tree(context)
    tree.append(dict(event))
    if on_event is not None:
        on_event(event)
        return
    if on_split is not None and event.get('event') == 'provider_split':
        # Legacy callback used by older tests: (exc, sources-like names only via event).
        class _Exc(Exception):
            pass
        legacy = _Exc(event.get('error') or event.get('reason') or 'provider_split')
        if event.get('http_status') is not None:
            legacy.status_code = event['http_status']
        sources = [
            type('Source', (), {'name': name})()
            for name in (event.get('filenames') or [])
        ]
        on_split(legacy, sources)


def inspect_resilient(
    inspect_once,
    sources,
    context,
    on_split=None,
    *,
    on_event=None,
    depth=0,
    parent_id=None,
    node_id=None,
    branch=None,
):
    """Inspect authorized images. Provider HTTP 500s split into smaller pieces.

    Quality settings are unchanged. Split depth and minimum part size are capped
    so repeated failures cannot fan out without bound.
    """
    node_id = node_id or _split_node_id(context, depth, parent_id, branch)
    try:
        output, usage = inspect_once(sources, context)
        # Record billed leaves that exist because a parent split. A successful
        # first request (no parent) stays quiet so normal inspection is unchanged.
        if parent_id is not None:
            billing = _usage_billing_fields(usage if isinstance(usage, dict) else {})
            _emit_split_event(on_event, on_split, context, {
                'event': 'provider_split_leaf',
                'kind': 'successful_billed_leaf' if billing['billed'] is True else 'successful_leaf',
                'node_id': node_id,
                'parent_id': parent_id,
                'depth': depth,
                'reason': 'child_after_parent_provider_failure',
                'http_status': 200,
                'source_image_count': len(sources),
                'filenames': [getattr(source, 'name', '') for source in sources],
                'application_retry': False,
                'provider_api_retry': False,
                'child_split_request': True,
                **billing,
            })
        return output, usage
    except Exception as exc:
        converted = as_credit_error(exc)
        if converted is not exc:
            raise converted from exc
        http_status = _provider_http_status(exc)
        filenames = [getattr(source, 'name', '') for source in sources]
        if not provider_server_error(exc):
            raise
        if len(sources) < PROVIDER_SPLIT_MIN_IMAGES:
            filename = filenames[0] if filenames else 'this image'
            _emit_split_event(on_event, on_split, context, {
                'event': 'provider_split_exhausted',
                'kind': 'minimum_part_size',
                'node_id': node_id,
                'parent_id': parent_id,
                'depth': depth,
                'reason': 'part_below_minimum_split_size',
                'http_status': http_status,
                'source_image_count': len(sources),
                'filenames': filenames,
                'billed': False,
                'billing_status': 'unbilled_http_failure',
                'response_id': None,
                'token_usage': None,
                'application_retry': False,
                'provider_api_retry': False,
                'child_split_request': False,
                'error_type': type(exc).__name__,
                'error': str(exc)[:500],
            })
            raise RuntimeError(
                f'OpenAI rejected {filename} with a server error. The other images in the batch can be inspected separately.'
            ) from exc
        if depth >= PROVIDER_SPLIT_MAX_DEPTH:
            _emit_split_event(on_event, on_split, context, {
                'event': 'provider_split_exhausted',
                'kind': 'max_depth',
                'node_id': node_id,
                'parent_id': parent_id,
                'depth': depth,
                'reason': 'provider_split_max_depth_reached',
                'http_status': http_status,
                'source_image_count': len(sources),
                'filenames': filenames,
                'billed': False,
                'billing_status': 'unbilled_http_failure',
                'response_id': None,
                'token_usage': None,
                'application_retry': False,
                'provider_api_retry': False,
                'child_split_request': False,
                'max_depth': PROVIDER_SPLIT_MAX_DEPTH,
                'error_type': type(exc).__name__,
                'error': str(exc)[:500],
            })
            raise ProviderSplitExhausted(
                f'Provider split depth limit ({PROVIDER_SPLIT_MAX_DEPTH}) reached for '
                f'{len(sources)} images starting at {filenames[0] if filenames else "batch"}. '
                'No further child requests were created. Inspection stopped without staging.'
            ) from exc
        _emit_split_event(on_event, on_split, context, {
            'event': 'provider_split',
            'kind': 'unsuccessful_parent_request',
            'node_id': node_id,
            'parent_id': parent_id,
            'depth': depth,
            'reason': 'provider_server_error',
            'http_status': http_status,
            'source_image_count': len(sources),
            'filenames': filenames,
            'billed': False,
            'billing_status': 'unbilled_http_failure',
            'response_id': None,
            'token_usage': None,
            'application_retry': False,
            'provider_api_retry': False,
            'child_split_request': True,
            'error_type': type(exc).__name__,
            'error': str(exc)[:500],
        })
        midpoint = len(sources) // 2
        left_output, left_usage = inspect_resilient(
            inspect_once, sources[:midpoint], context, on_split,
            on_event=on_event, depth=depth + 1, parent_id=node_id, branch='L',
        )
        right_context = continuation_context(context, left_output)
        right_output, right_usage = inspect_resilient(
            inspect_once, sources[midpoint:], right_context, on_split,
            on_event=on_event, depth=depth + 1, parent_id=node_id, branch='R',
        )
        continued = right_context.get('prior_open_item')
        return merge_inspections([left_output, right_output], [continued]), merge_usage([left_usage, right_usage])

def _runtime_prompt(settings):
    configured = getattr(settings, 'runtime_prompt_path', None) or ''
    candidates = []
    if configured:
        candidates.append(Path(configured))
    app_dir = Path(__file__).resolve().parents[1]
    candidates.append(app_dir.parent / 'Runtime_prompt_old_chatgpt.md')
    candidates.append(app_dir / 'prompts' / 'Runtime_prompt_old_chatgpt.md')
    for path in candidates:
        if path.is_file():
            return path.read_text(encoding='utf-8'), path.name
    raise ValueError('Canonical runtime prompt is missing. Set RUNTIME_PROMPT_PATH to Runtime_prompt_old_chatgpt.md')

def _docx_text(path):
    from docx import Document
    document = Document(path)
    parts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    for index, table in enumerate(document.tables, 1):
        rows = []
        for row in table.rows:
            cells = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
            if any(cells):
                rows.append(' | '.join(cells))
        if rows:
            parts.append(f'Table {index}\n' + '\n'.join(rows))
    return '\n'.join(parts)

class Inspector:
    def __init__(self, settings, drive, client=None):
        self.settings = settings
        self.drive = drive
        self.diagnostics = Diagnostics(getattr(settings, 'diagnostics_dir', './logs'))
        if client is None:
            from openai import OpenAI
            client = OpenAI(api_key=settings.openai_api_key, max_retries=0, timeout=getattr(settings, 'openai_timeout_seconds', 900))
        self.client = client
        self.usage_recorder = None
        self._sleep = time.sleep

    def _call(self, job, batch, attempt, operation, func):
        delay = 2
        for try_index in range(1, CONNECTION_ATTEMPTS + 1):
            try:
                return func()
            except Exception as exc:
                timed_out = _is_ambiguous_timeout(exc)
                if timed_out:
                    # A timeout can happen after the provider accepted the request.
                    # Do not call that try unbilled, and do not invent token counts.
                    self.diagnostics.write(
                        job, batch, attempt, 'usage_outcome_unknown',
                        operation=operation, try_index=try_index, attempts=CONNECTION_ATTEMPTS,
                        billed=None, outcome='unknown',
                        failure_reason='timeout_may_have_been_processed',
                        error_type=type(exc).__name__,
                    )
                if not _is_connection_error(exc):
                    raise
                if try_index != CONNECTION_ATTEMPTS:
                    self.diagnostics.write(
                        job, batch, attempt, 'connection_retry',
                        operation=operation, try_index=try_index, attempts=CONNECTION_ATTEMPTS,
                        error_type=type(exc).__name__, error=str(exc), error_chain=exception_chain(exc),
                    )
                if not timed_out:
                    # The connection failed before a response. Do not invent token counts.
                    self.diagnostics.write(
                        job, batch, attempt, 'usage_unbilled',
                        operation=operation, try_index=try_index, attempts=CONNECTION_ATTEMPTS,
                        billed=False, failure_reason=str(exc)[:300], error_type=type(exc).__name__,
                    )
                if try_index == CONNECTION_ATTEMPTS:
                    raise
                self._sleep(delay)
                delay = min(delay * 2, 16)

    def _download_drive_image(self, job, batch, attempt, source):
        """Retry only Drive media. Never retries inspection or OpenAI calls."""
        filename = getattr(source, 'name', '') or getattr(source, 'id', 'image')
        delay = DRIVE_DOWNLOAD_BACKOFF_SECONDS
        for try_index in range(1, DRIVE_DOWNLOAD_ATTEMPTS + 1):
            try:
                raw = self.drive.image(source)
                if hasattr(raw, 'read'):
                    raw = raw.read()
                if not isinstance(raw, (bytes, bytearray)):
                    raise TypeError(f'Drive image for {filename} did not return bytes')
                raw = bytes(raw)
                if not raw:
                    raise IncompleteDriveDownload(f'{filename}: downloaded 0 bytes')
                return raw
            except Exception as exc:
                if not _is_drive_download_error(exc):
                    if filename and filename not in str(exc):
                        raise RuntimeError(f'Drive download failed for {filename}: {exc}') from exc
                    raise
                try:
                    self.diagnostics.write(
                        job, batch, attempt, 'drive_download_retry',
                        filename=filename,
                        try_index=try_index,
                        attempts=DRIVE_DOWNLOAD_ATTEMPTS,
                        error_type=type(exc).__name__,
                        error=str(exc)[:500],
                        error_chain=exception_chain(exc),
                    )
                except Exception:
                    pass
                if try_index == DRIVE_DOWNLOAD_ATTEMPTS:
                    raise RuntimeError(
                        f'Drive download failed for {filename} after {DRIVE_DOWNLOAD_ATTEMPTS} attempts: '
                        f'{type(exc).__name__}: {exc}'
                    ) from exc
                self._sleep(delay)
                delay = min(delay * 2, 4)

    def _image_content(self, job, batch, attempt, blobs, file_ids):
        files_api = getattr(getattr(self.client, 'files', None), 'create', None)
        content = []
        for src, data, mime in blobs:
            content.append({'type': 'input_text', 'text': f'Source filename: {src.name}; Drive ID: {src.id}'})
            if files_api is None:
                content.append({
                    'type': 'input_image',
                    'image_url': f'data:{mime};base64,{base64.b64encode(data).decode()}',
                    'detail': 'high',
                })
                continue
            uploaded = self._call(
                job, batch, attempt, f'files.create {src.name}',
                lambda data=data, mime=mime, name=src.name: self.client.files.create(
                    file=(name, data, mime),
                    purpose='user_data',
                    expires_after={'anchor': 'created_at', 'seconds': 3600},
                ),
            )
            file_id = getattr(uploaded, 'id', '') or ''
            if not file_id:
                raise ValueError(f'Image upload did not return a file id for {src.name}')
            file_ids.append(file_id)
            content.append({'type': 'input_image', 'file_id': file_id, 'detail': 'high'})
        return content

    def _delete_files(self, file_ids):
        delete = getattr(getattr(self.client, 'files', None), 'delete', None)
        if delete is None:
            return
        for file_id in file_ids:
            try:
                delete(file_id)
            except Exception:
                pass

    def inspect(self, job, batch, sources, context, rules):
        def once(chosen, chosen_context):
            return self._inspect_once(job, batch, chosen, chosen_context, rules)

        attempt = int((context or {}).get('inspection_attempt') or 0) if isinstance(context, dict) else 0

        def on_event(event):
            try:
                name = str(event.get('event') or 'provider_split')
                payload = {key: value for key, value in event.items() if key != 'event'}
                self.diagnostics.write(job, batch, attempt, name, **payload)
            except Exception:
                pass

        def _write_tree(reason: str | None = None):
            tree = _split_tree(context) if isinstance(context, dict) else []
            if not tree:
                return None
            try:
                return self.diagnostics.write(
                    job, batch, attempt, 'provider_split_tree',
                    tree=list(tree),
                    max_depth=PROVIDER_SPLIT_MAX_DEPTH,
                    min_images=PROVIDER_SPLIT_MIN_IMAGES,
                    reason=reason,
                )
            except Exception:
                return None

        try:
            result = inspect_resilient(once, sources, context, on_event=on_event)
        except ProviderSplitExhausted as exc:
            path = _write_tree(str(exc)[:500])
            if path is not None:
                setattr(exc, 'diagnostics_path', str(path))
            setattr(exc, 'split_tree', list(_split_tree(context)) if isinstance(context, dict) else [])
            raise
        except Exception:
            _write_tree()
            raise
        _write_tree()
        return result


    def _inspect_once(self, job, batch, sources, context, rules):
        prompt_text, prompt_name = _runtime_prompt(self.settings)
        docs = Path(self.settings.governing_docs_dir)
        references = []
        doc_names = []
        for path in sorted(docs.iterdir()):
            if path.suffix.lower() in {'.md', '.txt'}:
                references.append(path.name + '\n' + path.read_text(encoding='utf-8'))
                doc_names.append(path.name)
            elif path.suffix.lower() == '.docx':
                references.append(path.name + '\n' + _docx_text(path))
                doc_names.append(path.name)
        if not references:
            raise ValueError('No readable governing documents')
        approved = context.get('approved_project_reference') if isinstance(context, dict) else None
        reference_id = ''
        reference_version = ''
        if isinstance(approved, dict):
            reference_id = approved.get('reference_id') or approved.get('source_file') or ''
            reference_version = str(approved.get('version') or '')
        blobs = []
        modes = []
        attempt = int(getattr(context, 'attempt', 0) or 0) if not isinstance(context, dict) else int(context.get('inspection_attempt', 0) or 0)
        start = time.monotonic()
        original_max_bytes = int(getattr(self.settings, 'image_original_max_bytes', ORIGINAL_MAX_BYTES) or ORIGINAL_MAX_BYTES)
        original_max_edge = int(getattr(self.settings, 'image_original_max_edge', ORIGINAL_MAX_EDGE) or ORIGINAL_MAX_EDGE)
        resize_max_edge = int(getattr(self.settings, 'image_resize_max_edge', RESIZE_MAX_EDGE) or RESIZE_MAX_EDGE)
        jpeg_quality = int(getattr(self.settings, 'image_jpeg_quality', JPEG_QUALITY) or JPEG_QUALITY)
        try:
            for src in sources:
                raw = self._download_drive_image(job, batch, attempt, src)
                source_bytes = len(raw)
                data, mime, mode = prepare_model_image(
                    raw,
                    getattr(src, 'mime_type', '') or 'image/jpeg',
                    original_max_bytes=original_max_bytes,
                    original_max_edge=original_max_edge,
                    resize_max_edge=resize_max_edge,
                    jpeg_quality=jpeg_quality,
                )
                modes.append({'filename': src.name, 'mode': mode, 'mime': mime, 'source_bytes': source_bytes, 'prepared_bytes': len(data)})
                blobs.append((src, data, mime))
        except Exception as exc:
            setattr(exc, 'elapsed_seconds', round(time.monotonic() - start, 2))
            failing = ''
            if modes:
                prepared = {item.get('filename') for item in modes}
                remaining = [getattr(src, 'name', '') for src in sources if getattr(src, 'name', '') not in prepared]
                failing = remaining[0] if remaining else ''
            try:
                path = self.diagnostics.write(
                    job, batch, attempt, 'inspection_error', error_type=type(exc).__name__, error=str(exc),
                    error_chain=exception_chain(exc),
                    elapsed_seconds=getattr(exc, 'elapsed_seconds', 0), success=False, image_preparation=modes,
                    image_count=len(modes),
                    source_bytes_total=sum(item.get('source_bytes') or 0 for item in modes),
                    prepared_bytes_total=sum(item.get('prepared_bytes') or 0 for item in modes),
                    failing_filename=failing or None,
                )
                setattr(exc, 'diagnostics_path', str(path))
            except Exception:
                pass
            raise
        instruction = (
            'The client runtime prompt below is the canonical inspection workflow. '
            'Return Lean Split v2 Asset Staging and Record Staging in the structured object. '
            'Do not update Google Sheets. '
            'A client-approved calibration example in the user payload is calibration for the named project and box, not a filename rule and not an answer-key extract. '
            'Inspect every authorized image. Do not copy an example when the images disagree, and do not close an item only because its folder has ended.\n\n'
            + prompt_text
            + '\n\n'
            + '\n\n'.join(references)
            + '\n\n'
            + OUTPUT_SCHEMA
            + '\n\n'
            + decision_instructions()
        )
        prior = context.get('prior_open_item') if isinstance(context, dict) else None
        prior_text = 'None' if prior in (None, '', 'None') else str(prior)
        existing = []
        if isinstance(context, dict) and prior_text != 'None':
            existing = context.get('existing_open_record') or context.get('prior_open_item_context') or []
        payload = {
            'job_id': job,
            'batch_id': batch,
            'project': context.get('project') if isinstance(context, dict) else '',
            'run_context': context.get('run_context') if isinstance(context, dict) else [],
            'batch_queue_row': context.get('batch_queue_row') if isinstance(context, dict) else {},
            'manifest': context.get('manifest') if isinstance(context, dict) else [],
            'prior_open_item': prior_text,
            'prior_open_item_source': (context.get('prior_open_item_source') if isinstance(context, dict) else '') or 'Batch Queue Prior Open Item',
            'existing_open_record': existing,
            'prior_open_item_context': existing,
            'project_classification_rules': rules,
            'authorized_project_state': context.get('authorized_project_state') if isinstance(context, dict) else {},
            'approved_project_reference': approved,
            'folder_code': context.get('folder_code') if isinstance(context, dict) else '',
            'authorized_filenames': [s.name for s in sources],
        }
        start = time.monotonic()
        response_meta = {}
        file_ids = []
        attempt = int(getattr(context, 'attempt', 0) or 0) if not isinstance(context, dict) else int(context.get('inspection_attempt', 0) or 0)
        max_tokens = getattr(self.settings, 'openai_max_output_tokens', 0)
        effort = (getattr(self.settings, 'openai_reasoning_effort', '') or '').strip()
        prompt_ids = [prompt_name, *doc_names]
        if reference_id:
            prompt_ids.append(reference_id)
        try:
            images = self._image_content(job, batch, attempt, blobs, file_ids)
            image_transport = 'openai_files' if file_ids else 'inline'
            self.diagnostics.write(
                job, batch, attempt, 'request_prepared',
                model=self.settings.openai_model,
                reasoning_effort=effort,
                max_output_tokens=max_tokens,
                source_image_count=len(sources),
                authorized_filenames=[s.name for s in sources],
                prior_open_item=prior_text,
                prior_open_item_source=payload['prior_open_item_source'],
                prompt_identifiers=prompt_ids,
                reference_id=reference_id,
                reference_version=reference_version,
                image_preparation=modes,
                image_count=len(modes),
                image_transport=image_transport,
                uploaded_file_count=len(file_ids),
                source_bytes_total=sum(item.get('source_bytes') or 0 for item in modes),
                prepared_bytes_total=sum(item.get('prepared_bytes') or 0 for item in modes),
                prompt_character_count=len(instruction),
                rules_count=len(rules),
            )
            request = {
                'model': self.settings.openai_model,
                'input': [
                    {'role': 'system', 'content': instruction},
                    {'role': 'user', 'content': [{'type': 'input_text', 'text': json.dumps(payload)}] + images},
                ],
                'text_format': InspectionOutput,
            }
            if max_tokens:
                request['max_output_tokens'] = max_tokens
            if effort:
                request['reasoning'] = {'effort': effort}
            response = self._call(job, batch, attempt, 'responses.parse', lambda: self.client.responses.parse(**request))
            from app.usage_accounting import parse_provider_usage
            part = _next_part(context if isinstance(context, dict) else {})
            response_meta = parse_provider_usage(
                getattr(response, 'usage', None),
                response_id=getattr(response, 'id', '') or '',
                attempt=attempt,
                part=part,
                status='Received',
            )
            response_meta['provider_status'] = getattr(response, 'status', '')
            response_meta['incomplete_details'] = getattr(response, 'incomplete_details', None)
            response_meta['job_id'] = job
            response_meta['batch_id'] = batch
            response_meta['elapsed_seconds'] = round(time.monotonic() - start, 2)
            output = response.output_parsed
            self.diagnostics.write(
                job, batch, attempt, 'response_received', **{
                    key: value for key, value in response_meta.items()
                    if key not in {'incomplete_details', 'attempt'}
                },
                returned_asset_filenames=[a.digital_asset for a in output.assets] if output else [],
                returned_record_ids=[r.proposed_record_id for r in output.records] if output else [],
            )
            if self.usage_recorder is not None:
                stored = self.usage_recorder(response_meta)
                response_meta['usage_recorded'] = True
                if isinstance(stored, dict) and stored.get('estimated_cost_usd') not in (None, ''):
                    response_meta['estimated_cost_usd'] = stored['estimated_cost_usd']
            if output is None:
                raise ValueError('Missing structured output')
            expected = [s.name for s in sources]
            actual = [a.digital_asset for a in output.assets]
            if actual != expected:
                missing = sorted(set(expected) - set(actual))
                extra = sorted(set(actual) - set(expected))
                self.diagnostics.write(
                    job, batch, attempt, 'inspection_validation_failed',
                    reason='filename_integrity', expected_count=len(expected),
                    returned_count=len(actual), missing=missing, extra=extra,
                    expected_order=expected, returned_order=actual,
                )
                raise ValueError(
                    f'Inspection returned incomplete/non-ordered staging: expected {len(expected)} assets, got {len(actual)}; missing={missing[:5]}; extra={extra[:5]}'
                )
            if len(output.records) == 0:
                self.diagnostics.write(job, batch, attempt, 'inspection_validation_failed', reason='no_record_rows')
                raise ValueError('Inspection returned no Record Staging rows')
            linked = {a.assigned_record_id for a in output.assets}
            proposed = {r.proposed_record_id for r in output.records}
            if not all(a.assigned_record_id for a in output.assets) or not linked.issubset(proposed):
                self.diagnostics.write(
                    job, batch, attempt, 'inspection_validation_failed', reason='unlinked_assets',
                    linked_record_ids=sorted(linked), returned_record_ids=sorted(proposed),
                )
                raise ValueError('Every inspected asset must link to a returned Record Staging row')
            elapsed = response_meta['elapsed_seconds']
            self.diagnostics.write(
                job, batch, attempt, 'inspection_validated',
                elapsed_seconds=elapsed,
                success=True,
                image_preparation=modes,
                estimated_cost_usd=response_meta.get('estimated_cost_usd'),
                response_id=response_meta['response_id'],
            )
            usage_out = {
                key: response_meta.get(key)
                for key in (
                    'input_tokens', 'uncached_input_tokens', 'cached_input_tokens',
                    'cache_write_30m_tokens', 'cache_write_12h_tokens', 'output_tokens',
                    'response_id', 'attempt', 'part', 'status', 'failure_reason',
                    'usage_identity_gap', 'billed', 'usage_recorded', 'estimated_cost_usd',
                )
            }
            usage_out['elapsed_seconds'] = elapsed
            usage_out['responses'] = [{key: value for key, value in usage_out.items() if key != 'responses'}]
            return output, usage_out
        except Exception as e:
            setattr(e, 'elapsed_seconds', round(time.monotonic() - start, 2))
            if response_meta:
                carried = {
                    key: response_meta.get(key)
                    for key in (
                        'input_tokens', 'uncached_input_tokens', 'cached_input_tokens',
                        'cache_write_30m_tokens', 'cache_write_12h_tokens', 'output_tokens',
                        'response_id', 'attempt', 'part', 'status', 'failure_reason',
                        'usage_identity_gap', 'billed', 'usage_recorded', 'estimated_cost_usd',
                    )
                }
                carried['elapsed_seconds'] = getattr(e, 'elapsed_seconds', 0)
                setattr(e, 'inspection_usage', carried)
            try:
                path = self.diagnostics.write(
                    job, batch, attempt, 'inspection_error',
                    error_type=type(e).__name__,
                    error=str(e),
                    error_chain=exception_chain(e),
                    elapsed_seconds=getattr(e, 'elapsed_seconds', 0),
                    success=False,
                    image_preparation=modes,
                    image_count=len(modes),
                    source_bytes_total=sum(item.get('source_bytes') or 0 for item in modes),
                    prepared_bytes_total=sum(item.get('prepared_bytes') or 0 for item in modes),
                    prior_open_item_source=payload['prior_open_item_source'],
                    prompt_identifiers=prompt_ids,
                    estimated_cost_usd=(response_meta or {}).get('estimated_cost_usd'),
                    response_id=(response_meta or {}).get('response_id', ''),
                )
                setattr(e, 'diagnostics_path', str(path))
            except Exception:
                pass
            credit = as_credit_error(e)
            if credit is not e:
                raise credit from e
            raise
        finally:
            self._delete_files(file_ids)
